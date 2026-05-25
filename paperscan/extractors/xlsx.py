"""XLSX extractor — covers spreadsheet injection (hidden rows, hidden sheets, comments)."""
from __future__ import annotations

from paperscan.models import ExtractedDocument


def extract_xlsx(path: str) -> ExtractedDocument:
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required for XLSX support: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)

    # Workbook-level metadata
    metadata: dict = {}
    props = wb.properties
    if props:
        for attr in ("title", "subject", "description", "creator", "keywords", "category"):
            val = getattr(props, attr, None)
            if val:
                metadata[attr] = str(val)
    metadata["sheets"] = ", ".join(wb.sheetnames)

    visible_parts: list[str] = []
    hidden_text: list[dict] = []
    annotations: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_hidden = ws.sheet_state != "visible"
        sheet_text: list[str] = []

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                val = str(cell.value).strip()
                if not val:
                    continue

                row_dim = ws.row_dimensions.get(cell.row)
                col_dim = ws.column_dimensions.get(cell.column_letter)
                row_hidden = bool(row_dim and row_dim.hidden)
                col_hidden = bool(col_dim and col_dim.hidden)

                if sheet_hidden or row_hidden or col_hidden:
                    method = (
                        "hidden_sheet" if sheet_hidden
                        else "hidden_row" if row_hidden
                        else "hidden_column"
                    )
                    hidden_text.append({
                        "location": f"sheet:{sheet_name}!{cell.coordinate}",
                        "content": val,
                        "method": method,
                    })
                else:
                    sheet_text.append(val)

                # Cell comments
                if cell.comment:
                    comment_text = str(cell.comment.text or "").strip()
                    if comment_text:
                        annotations.append(
                            f"[comment@{sheet_name}!{cell.coordinate}]: {comment_text}"
                        )

        if sheet_text and not sheet_hidden:
            visible_parts.append(f"--- Sheet: {sheet_name} ---\n" + "  ".join(sheet_text))

    return ExtractedDocument(
        visible_text="\n\n".join(visible_parts),
        hidden_text=hidden_text,
        metadata=metadata,
        annotations=annotations,
    )
