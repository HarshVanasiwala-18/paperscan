"""PowerPoint (.pptx) extractor."""
from __future__ import annotations

import zipfile

from paperscan.models import ExtractedDocument

_MAX_ZIP_ENTRY = 100 * 1024 * 1024  # 100 MB per entry


def _check_pptx_zip(path: str) -> None:
    """Reject decompression bombs before passing the file to python-pptx."""
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            if info.file_size > _MAX_ZIP_ENTRY:
                raise ValueError(
                    f"PPTX entry '{info.filename}' exceeds size limit ({info.file_size} bytes)"
                )


def extract_pptx(path: str) -> ExtractedDocument:
    _check_pptx_zip(path)
    from pptx import Presentation  # type: ignore[import-untyped]

    prs = Presentation(path)
    visible_parts: list[str] = []
    hidden_text: list[dict] = []
    annotations: list[str] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_label = f"slide{slide_num}"

        for shape in slide.shapes:
            # Visible text frames
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        visible_parts.append(text)

            # Grouped shapes
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                for child in shape.shapes:
                    if child.has_text_frame:
                        for para in child.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                visible_parts.append(text)

            # QR codes in picture shapes
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    from paperscan.extractors.qr import scan_image_for_qr
                    for qr_text in scan_image_for_qr(shape.image.blob):
                        hidden_text.append({
                            "location": f"slide{slide_num}_image",
                            "content": qr_text,
                            "method": "qr_code",
                        })
                except Exception:
                    pass

        # Speaker notes — hidden from audience, readable by AI
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    hidden_text.append({
                        "location": f"{slide_label}_notes",
                        "content": notes,
                        "method": "speaker_notes",
                    })
        except Exception:
            pass

    # Core properties (title, author, subject, keywords, description)
    try:
        props = prs.core_properties
        metadata: dict = {}
        for attr in ("author", "title", "subject", "keywords", "description", "last_modified_by"):
            val = getattr(props, attr, None)
            if val:
                metadata[attr] = str(val)
    except Exception:
        metadata = {}

    return ExtractedDocument(
        visible_text="\n".join(visible_parts),
        hidden_text=hidden_text,
        annotations=annotations,
        metadata=metadata,
    )
